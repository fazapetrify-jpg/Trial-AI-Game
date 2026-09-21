import sys
import unittest
import uuid
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

HERE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(HERE))

from book import build_docx
from author import generate_manuscript
from font_embed import FONTS, R, W, obfuscate
from qa import audit_style
from video import clean_captions


class GeneratorTests(unittest.TestCase):
    def test_caption_cleanup(self):
        raw = "WEBVTT\n\n00:00:01.000 --> 00:00:03.000\nHalo semua\n\n00:00:03.000 --> 00:00:05.000\nHalo semua\nselamat belajar"
        self.assertEqual(clean_captions(raw), "Halo semua\nselamat belajar")

    def test_sbv_cleanup(self):
        raw = "0:00:01.000,0:00:03.000\nKlef G\n\n0:00:03.000,0:00:05.000\nuntuk membaca nada"
        self.assertEqual(clean_captions(raw), "Klef G\nuntuk membaca nada")

    def test_rolling_captions_do_not_repeat_sentences(self):
        raw = "WEBVTT\n00:00:01.000 --> 00:00:02.000\nHalo semua\n00:00:02.000 --> 00:00:03.000\nHalo semua selamat belajar\n00:00:03.000 --> 00:00:04.000\nHalo semua selamat belajar\n"
        self.assertEqual(clean_captions(raw), "Halo semua\nselamat belajar")

    def test_transcript_authoring_pipeline(self):
        import json
        from types import SimpleNamespace

        class FakeResponses:
            def __init__(self):
                self.calls = []

            def create(self, **kwargs):
                self.calls.append(kwargs)
                outputs = [
                    {"notes": ["Klef G menentukan acuan nada pada garis paranada."]},
                    {"chapters": [{"title": "Membaca paranada", "focus": "Klef G"}],
                     "introduction": ["Mulai dari garis dan spasi."]},
                    {"sections": [{"heading": "Klef G", "paragraphs": ["Kenali posisi nada terlebih dahulu."],
                                   "bullets": [], "tip": "Baca perlahan."}]},
                ]
                return SimpleNamespace(output_text=json.dumps(outputs[len(self.calls) - 1]))

        fake = FakeResponses()
        manuscript = generate_manuscript("Klef G dan paranada. " * 15, title="Belajar Not Balok",
                                        client=SimpleNamespace(responses=fake))
        self.assertEqual(manuscript["chapters"][0]["sections"][0]["heading"], "Klef G")
        self.assertEqual(len(fake.calls), 3)
        self.assertTrue(all(call["store"] is False for call in fake.calls))

    def test_font_obfuscation_is_reversible(self):
        data = b"\x00\x01\x00\x00" + bytes(range(4, 64))
        key = uuid.UUID("001B70DC-AA60-4AD5-90EC-18A0948E1EAE")
        self.assertEqual(obfuscate(obfuscate(data, key), key), data)

    def test_build_matches_reference_typography(self):
        import json
        from docx import Document
        manuscript = json.loads((HERE / "example_manuscript.json").read_text(encoding="utf-8"))
        output = HERE / "tests" / "sample_test.docx"
        build_docx(manuscript, output, base=HERE)
        self.assertEqual(audit_style(output), [])
        doc = Document(output)
        self.assertAlmostEqual(doc.sections[0].page_width.inches, 5.827, places=2)
        self.assertAlmostEqual(doc.sections[0].top_margin.inches, 1.0, places=2)
        self.assertEqual(doc.styles["Book Section"].font.size.pt, 14)
        self.assertEqual(doc.styles["Book Subheading"].font.size.pt, 12)
        self.assertEqual(str(doc.styles["Heading 1"].font.color.rgb), "000000")
        with zipfile.ZipFile(output) as archive:
            fonts = [x for x in archive.namelist() if x.endswith(".odttf")]
            self.assertEqual(len(fonts), 2)
            table = ET.fromstring(archive.read("word/fontTable.xml"))
            self.assertIn("Montserrat", ET.tostring(table, encoding="unicode"))
            self.assertIsNotNone(ET.fromstring(archive.read("word/settings.xml")).find(f"{{{W}}}embedTrueTypeFonts"))
            font = next(x for x in table if x.get(f"{{{W}}}name") == "Montserrat")
            relationships = ET.fromstring(archive.read("word/_rels/fontTable.xml.rels"))
            for name, face in (("embedRegular", "Regular"), ("embedBold", "Bold")):
                node = font.find(f"{{{W}}}{name}")
                guid = uuid.UUID(node.get(f"{{{W}}}fontKey"))
                rid = node.get(f"{{{R}}}id")
                target = next(x.get("Target") for x in relationships if x.get("Id") == rid)
                recovered = obfuscate(archive.read("word/" + target), guid)
                self.assertEqual(recovered, (FONTS / f"Montserrat-{face}.ttf").read_bytes())

    def test_ordered_layout_blocks(self):
        from docx import Document
        manuscript = {
            "title": "Contoh", "chapters": [{"title": "Dasar", "sections": [{
                "heading": "Urutan", "blocks": [
                    {"type": "paragraph", "text": "Penjelasan satu."},
                    {"type": "subheading", "text": "Contoh"},
                    {"type": "paragraph", "text": "Penjelasan dua."},
                ],
            }]}],
        }
        output = HERE / "tests" / "sample_test.docx"
        build_docx(manuscript, output)
        text = [p.text for p in Document(output).paragraphs]
        self.assertLess(text.index("Penjelasan satu."), text.index("Contoh"))
        self.assertLess(text.index("Contoh"), text.index("Penjelasan dua."))


if __name__ == "__main__":
    unittest.main()
