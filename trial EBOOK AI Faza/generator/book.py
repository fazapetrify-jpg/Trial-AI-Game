"""Render an editable A5 DOCX using the measured Gaya Ngiring style profile."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image
from font_embed import embed_montserrat

ROOT = Path(__file__).resolve().parent
STYLE_PATH = ROOT / "style_gaya_ngiring.json"


def load_style(path: Path = STYLE_PATH) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def validate_manuscript(book: dict[str, Any]) -> None:
    if not isinstance(book, dict):
        raise ValueError("Manuscript must be a JSON object")
    if not isinstance(book.get("title"), str) or not book["title"].strip():
        raise ValueError("Manuscript needs a nonempty title")
    chapters = book.get("chapters")
    if not isinstance(chapters, list) or not chapters:
        raise ValueError("Manuscript needs at least one chapter")
    for number, chapter in enumerate(chapters, 1):
        if not isinstance(chapter, dict) or not isinstance(chapter.get("title"), str):
            raise ValueError(f"Chapter {number} needs a title")
        if not isinstance(chapter.get("sections"), list) or not chapter["sections"]:
            raise ValueError(f"Chapter {number} needs sections")
        for section in chapter["sections"]:
            if not isinstance(section, dict) or not isinstance(section.get("heading"), str):
                raise ValueError(f"A section in chapter {number} needs a heading")
            if not isinstance(section.get("paragraphs"), list) and not isinstance(section.get("blocks"), list):
                raise ValueError(f"Section {section['heading']} needs paragraphs or blocks")


def _set_font(style, family: str, size: float, color: str, bold: bool = False) -> None:
    style.font.name = family
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is not None:
        for part in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{part}"), family)


def _style_doc(doc: Document, spec: dict[str, Any]) -> None:
    section = doc.sections[0]
    section.page_width = Inches(spec["page_width_inches"])
    section.page_height = Inches(spec["page_height_inches"])
    section.top_margin = section.bottom_margin = Inches(spec["margin_inches"])
    section.left_margin = section.right_margin = Inches(spec["margin_inches"])
    section.header_distance = section.footer_distance = Inches(0.5)

    styles = doc.styles
    normal = styles["Normal"]
    _set_font(normal, spec["font_family"], spec["body_pt"], spec["text_color"])
    normal.paragraph_format.line_spacing = spec["body_line_spacing"]
    normal.paragraph_format.space_after = Pt(spec["body_after_pt"])
    normal.paragraph_format.first_line_indent = Pt(spec["body_first_indent_pt"])

    roles = [
        ("Book Section", spec["section_heading_pt"], True),
        ("Book Subheading", spec["subheading_pt"], True),
        ("Chapter Number", spec["chapter_number_pt"], True),
        ("Chapter Title", spec["chapter_title_pt"], False),
        ("Book Caption", 10, False),
    ]
    for name, size, bold in roles:
        style = styles[name] if name in styles else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
        _set_font(style, spec["font_family"], size, spec["heading_color"], bold)
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.line_spacing = spec["body_line_spacing"]
        style.paragraph_format.keep_with_next = name != "Book Caption"
        if name == "Book Section":
            style.paragraph_format.space_before = Pt(12)
            style.paragraph_format.space_after = Pt(12)
        elif name == "Book Subheading":
            style.paragraph_format.space_before = Pt(12)
            style.paragraph_format.space_after = Pt(3)
        else:
            style.paragraph_format.space_after = Pt(6)

    # Word's built-in Heading 1/2 defaults are blue. Explicitly neutralize them,
    # including theme-color attributes that can override RGB in another editor.
    for name, size in (("Heading 1", spec["section_heading_pt"]),
                       ("Heading 2", spec["subheading_pt"]),
                       ("Title", spec["chapter_number_pt"])):
        style = styles[name]
        _set_font(style, spec["font_family"], size, spec["heading_color"], True)
        color = style.element.get_or_add_rPr().color
        if color is not None:
            color.attrib.pop(qn("w:themeColor"), None)
            color.attrib.pop(qn("w:themeTint"), None)
            color.attrib.pop(qn("w:themeShade"), None)


def _paragraph(doc: Document, text: str, *, center: bool = False, style: str | None = None):
    paragraph = doc.add_paragraph(style=style)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.add_run(text)
    return paragraph


def _heading(doc: Document, text: str, style: str = "Book Section"):
    paragraph = _paragraph(doc, text, center=style in {"Book Section", "Chapter Number", "Chapter Title"}, style=style)
    paragraph.paragraph_format.first_line_indent = Pt(0)
    return paragraph


def _figure(doc: Document, figure: dict[str, Any], base: Path, max_width_inches: float):
    raw = figure.get("path")
    if not isinstance(raw, str) or not raw:
        raise ValueError("Figure needs a path")
    path = (base / raw).resolve()
    if not path.is_file():
        raise FileNotFoundError(path)
    with Image.open(path) as image:
        ratio = image.height / image.width
    width = min(max_width_inches, float(figure.get("width_inches", max_width_inches)))
    height = width * ratio
    max_height = float(figure.get("max_height_inches", 2.5))
    if height > max_height:
        height = max_height
        width = height / ratio
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.keep_with_next = bool(figure.get("caption"))
    p.add_run().add_picture(str(path), width=Inches(width), height=Inches(height))
    if figure.get("caption"):
        c = _paragraph(doc, str(figure["caption"]), center=True, style="Book Caption")
        c.runs[0].italic = True


def _section_blocks(section: dict[str, Any]) -> list[dict[str, Any]]:
    if "blocks" in section:
        blocks = section["blocks"]
        if not isinstance(blocks, list) or not blocks:
            raise ValueError(f"Section {section['heading']} needs nonempty blocks")
        return blocks
    blocks = [{"type": "paragraph", "text": text} for text in section["paragraphs"]]
    blocks += [{"type": "bullet", "text": text} for text in section.get("bullets", [])]
    figures = section.get("figures", [])
    if section.get("figure"):
        figures = [section["figure"], *figures]
    blocks += [{"type": "figure", **figure} for figure in figures]
    if section.get("tip"):
        blocks.append({"type": "tip", "text": section["tip"]})
    return blocks


def _render_block(doc: Document, block: dict[str, Any], base: Path, max_width: float) -> None:
    if not isinstance(block, dict):
        raise ValueError("Each section block must be an object")
    kind = block.get("type")
    if kind == "paragraph":
        _paragraph(doc, str(block.get("text", "")))
    elif kind == "subheading":
        _heading(doc, str(block.get("text", "")), "Book Subheading")
    elif kind == "bullet":
        p = _paragraph(doc, "•  " + str(block.get("text", "")))
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.left_indent = Pt(17)
    elif kind == "figure":
        _figure(doc, block, base, max_width)
    elif kind == "tip":
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.add_run("Tips Aplikasi : ").bold = True
        p.add_run(str(block.get("text", "")))
    else:
        raise ValueError(f"Unsupported section block type: {kind}")


def build_docx(book: dict[str, Any], output: Path, *, base: Path | None = None,
               style_path: Path = STYLE_PATH) -> Path:
    validate_manuscript(book)
    spec = load_style(style_path)
    output = Path(output)
    base = Path(base or ".").resolve()
    doc = Document()
    _style_doc(doc, spec)
    max_width = spec["page_width_inches"] - 2 * spec["margin_inches"]

    if book.get("author_note"):
        _heading(doc, "KATA PENULIS")
        for text in book["author_note"]:
            _paragraph(doc, str(text))
        if book.get("author"):
            p = _paragraph(doc, str(book["author"]))
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.first_line_indent = Pt(0)
        doc.add_page_break()

    if book.get("how_this_helps"):
        _heading(doc, "CARA BUKU INI BISA MEMBANTUMU")
        for text in book["how_this_helps"]:
            _paragraph(doc, str(text))
        doc.add_page_break()

    _heading(doc, "DAFTAR ISI")
    for i, chapter in enumerate(book["chapters"], 1):
        p = _paragraph(doc, f"BAB {i}  {chapter['title'].upper()}")
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.line_spacing = 2.0
    doc.add_page_break()

    if book.get("introduction"):
        _heading(doc, "PENDAHULUAN")
        for text in book["introduction"]:
            _paragraph(doc, str(text))
        doc.add_page_break()

    for i, chapter in enumerate(book["chapters"], 1):
        _heading(doc, f"BAB {i}", "Chapter Number")
        _heading(doc, chapter["title"], "Chapter Title")
        doc.add_page_break()
        for section in chapter["sections"]:
            _heading(doc, section["heading"], "Book Section")
            for block in _section_blocks(section):
                _render_block(doc, block, base, max_width)

    if book.get("closing"):
        doc.add_page_break()
        _heading(doc, "PENUTUP")
        for text in book["closing"]:
            _paragraph(doc, str(text))
    if book.get("sources"):
        _heading(doc, "SUMBER DAN VIDEO PENDAMPING", "Book Subheading")
        for item in book["sources"]:
            _paragraph(doc, str(item))

    doc.core_properties.title = book["title"]
    doc.core_properties.author = str(book.get("author", "Musti Musik"))
    doc.core_properties.subject = "E-book pembelajaran musik"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    embed_montserrat(output)
    return output
